import openpyxl
import pyautogui
import time


# ============================================================
# LOGIN
# ============================================================

def login(username, password):
    # --------------------------------------------------------
    # Step 1: Username
    # --------------------------------------------------------
    pyautogui.write(str(username), interval=0.03)
    pyautogui.press("enter")

    print(f"2. Username entered: {username}")

    # Give Google time to respond.
    time.sleep(5)

    # --------------------------------------------------------
    # Step 2: Manual verification checkpoint
    # --------------------------------------------------------
    print("\n" + "=" * 65)
    print("MANUAL CHECKPOINT")
    print("=" * 65)
    print()
    print("Look at Chrome.")
    print()
    print("If Google shows a CAPTCHA/verification:")
    print("  1. Complete the verification manually.")
    print("  2. Wait until the password page is displayed.")
    print("  3. Click inside the PASSWORD field.")
    print()
    print("If NO verification appeared:")
    print("  1. Simply make sure the PASSWORD field is ready.")
    print()
    print("Then return to this PowerShell window.")
    print()
    print("Press ENTER = continue with password")
    print("Type STOP   = stop the entire automation")
    print()
    print("=" * 65)

    decision = input("Your choice: ").strip().lower()

    if decision == "stop":
        print("\nAutomation stopped by user.")
        return False

    # Small delay after returning from PowerShell.
    time.sleep(2)

    # --------------------------------------------------------
    # Step 3: Password
    # --------------------------------------------------------
    pyautogui.write(str(password), interval=0.03)
    pyautogui.press("enter")

    print("3. Password entered")

    # Allow login to complete.
    time.sleep(5)

    # --------------------------------------------------------
    # Step 4: YouTube
    # --------------------------------------------------------
    pyautogui.hotkey("ctrl", "l")
    pyautogui.write("https://www.youtube.com", interval=0.03)
    pyautogui.press("enter")

    print("4. Entering into YouTube")

    time.sleep(10)

    return True


# ============================================================
# LOGOUT
# ============================================================

def logout():
    pyautogui.hotkey("ctrl", "l")

    logout_url = "https://accounts.google.com/logout"

    pyautogui.write(logout_url, interval=0.03)
    pyautogui.press("enter")

    print("5. Logged out successfully")

    time.sleep(5)


# ============================================================
# STARTUP
# ============================================================

print("\n" + "=" * 65)
print("LOGIN / LOGOUT AUTOMATION - V2")
print("=" * 65)

print("\nMake sure Chrome is open and focused.")

time.sleep(5)


# ============================================================
# OPEN GOOGLE LOGIN
# ============================================================

pyautogui.hotkey("ctrl", "l")

link = "https://accounts.google.com/signin"

pyautogui.write(link, interval=0.03)
pyautogui.press("enter")

print("1. Opening Google Sign-In")

time.sleep(5)


# ============================================================
# EXCEL
# ============================================================

xlsx_file_name = (
    r"E:\Office-Works\daily-random-codes"
    r"\Login-Logout local\cc.xlsx"
)

try:
    wb = openpyxl.load_workbook(xlsx_file_name)

    sheet_names = wb.sheetnames

    print("\nAvailable sheets:")

    for number, sheet_name in enumerate(sheet_names, start=1):
        print(f"{number}. {sheet_name}")

    # --------------------------------------------------------
    # Sheet selection
    # --------------------------------------------------------

    sheet_number = int(
        input("\nEnter the sheet number: ")
    )

    if sheet_number < 1 or sheet_number > len(sheet_names):
        raise ValueError("Invalid sheet number.")

    ws = wb.worksheets[sheet_number - 1]

    print(f"\nSelected sheet: {ws.title}")

    # --------------------------------------------------------
    # Row selection
    # --------------------------------------------------------

    rows_to_process = input(
        "Do you want to specify rows to process? (yes/no): "
    ).strip().lower()

    if rows_to_process == "yes":

        row_numbers = input(
            "Enter exact row numbers separated by commas "
            "(example: 1,6,10): "
        )

        rows = set()

        for value in row_numbers.split(","):
            value = value.strip()

            if value:
                row_number = int(value)

                if row_number >= 1:
                    rows.add(row_number)

        print(f"Selected rows: {sorted(rows)}")

    else:

        # You said you want selected rows only,
        # so don't accidentally process everything.
        print("\nNo rows selected.")
        print("Stopping safely.")
        raise SystemExit


    # ========================================================
    # PROCESS SELECTED ROWS
    # ========================================================

    all_rows = list(
        ws.iter_rows(values_only=True)
    )

    for i, row in enumerate(all_rows, start=1):

        # Skip rows that weren't selected.
        if i not in rows:
            continue

        # ----------------------------------------------------
        # Validate row
        # ----------------------------------------------------

        if len(row) < 3:
            print(f"\nRow {i} does not contain 3 columns.")
            print("Expected: SNO | Username | Password")
            continue

        sno, username, password = row

        if username is None or password is None:
            print(
                f"\nSkipping Row {i}: "
                "username/password is empty."
            )
            continue

        print("\n" + "=" * 65)
        print(f"Executing Row {i}")
        print(f"SNO: {sno}")
        print(f"Username: {username}")
        print("=" * 65)

        # ----------------------------------------------------
        # Login
        # ----------------------------------------------------

        success = login(username, password)

        # User requested STOP.
        if not success:
            print("\nAutomation stopped.")
            break

        # ----------------------------------------------------
        # Logout
        # ----------------------------------------------------

        logout()

        time.sleep(5)

        # ----------------------------------------------------
        # Prepare Google login for next selected account
        # ----------------------------------------------------

        if i != max(rows):

            pyautogui.hotkey("ctrl", "l")

            pyautogui.write(
                "https://accounts.google.com/signin",
                interval=0.03
            )

            pyautogui.press("enter")

            print("Opening Google Sign-In for next account")

            time.sleep(5)


    print("\n" + "=" * 65)
    print("Selected-row processing completed.")
    print("=" * 65)


except KeyboardInterrupt:

    print("\nAutomation interrupted by keyboard.")


except Exception as e:

    print(f"\nError: {e}")
